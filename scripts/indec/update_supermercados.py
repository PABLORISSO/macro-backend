"""
update_supermercados.py
Descarga la Encuesta de Supermercados del INDEC y genera:
  - backend/src/data/consumo/supermercados.csv         (Cuadro 1 — índices constantes)
  - backend/src/data/consumo/supermercados_canales.csv (Cuadro 4 — canal y medios de pago)
  - backend/src/data/consumo/supermercados_bocas.csv   (Cuadro 6 — ticket, m2, personal)
  - backend/src/data/consumo.json                      (summary con datos reales del último período)

Uso:
  python scripts/indec/update_supermercados.py
"""

import csv
import json
import io
from pathlib import Path
from datetime import date, datetime

import openpyxl
import requests

BACKEND_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = BACKEND_ROOT / "src" / "data" / "raw" / "indec" / "supermercados"
OUT_DIR = BACKEND_ROOT / "src" / "data" / "consumo"
OUT_CSV = OUT_DIR / "supermercados.csv"
OUT_CSV_CANALES = OUT_DIR / "supermercados_canales.csv"
OUT_CSV_BOCAS = OUT_DIR / "supermercados_bocas.csv"
OUT_JSON = BACKEND_ROOT / "src" / "data" / "consumo.json"

SOURCE_URL = "https://www.indec.gob.ar/ftp/cuadros/economia/serie_supermercados.xlsx"
RAW_FILE = RAW_DIR / "serie_supermercados.xlsx"

RAW_DIR.mkdir(parents=True, exist_ok=True)
OUT_DIR.mkdir(parents=True, exist_ok=True)


def to_float(v):
    try:
        return round(float(v), 4)
    except (TypeError, ValueError):
        return None


def to_millions(v):
    """Convierte miles de pesos → millones de pesos redondeado a 1 decimal."""
    try:
        return round(float(v) / 1000, 1)
    except (TypeError, ValueError):
        return None


def descargar():
    print(f"Descargando supermercados desde: {SOURCE_URL}")
    try:
        r = requests.get(SOURCE_URL, timeout=60)
        r.raise_for_status()
        RAW_FILE.write_bytes(r.content)
        print(f"✅ Descargado: {RAW_FILE}")
        return r.content
    except Exception as err:
        if RAW_FILE.exists():
            print(f"⚠️  No se pudo descargar ({err}), usando copia local.")
            return RAW_FILE.read_bytes()
        raise RuntimeError(f"No se pudo descargar y no hay archivo local en {RAW_FILE}") from err


def procesar_cuadro1(wb):
    """Índices a precios constantes — serie original, desestacionalizada y tendencia."""
    ws = wb["Cuadro 1"]
    rows_out = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if not isinstance(row[0], datetime):
            continue
        rows_out.append({
            "fecha": row[0].strftime("%Y-%m"),
            "indice_orig": to_float(row[1]),
            "var_yoy": to_float(row[2]),
            "var_acum": to_float(row[3]),
            "indice_desest": to_float(row[5]),
            "var_mom_desest": to_float(row[6]),
            "indice_tend": to_float(row[8]),
            "var_mom_tend": to_float(row[9]),
        })
    return rows_out


def procesar_cuadro4(wb):
    """Canal de venta y medios de pago — valores en millones de pesos corrientes."""
    # cols: 0=fecha, 1=total, 3=salon_ventas, 4=online, 6=efectivo, 7=debito, 8=credito, 9=otros
    ws = wb["Cuadro 4."]
    rows_out = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if not isinstance(row[0], datetime):
            continue
        rows_out.append({
            "fecha": row[0].strftime("%Y-%m"),
            "total_mill": to_millions(row[1]),
            "salon_ventas_mill": to_millions(row[3]),
            "online_mill": to_millions(row[4]),
            "efectivo_mill": to_millions(row[6]),
            "debito_mill": to_millions(row[7]),
            "credito_mill": to_millions(row[8]),
            "otros_mill": to_millions(row[9]),
        })
    return rows_out


def procesar_cuadro6(wb):
    """Bocas, superficie, ticket promedio, ventas/m², personal."""
    # cols: 0=fecha, 1=ventas_miles_pesos, 2=bocas, 3=ventas_por_boca, 4=m2, 5=ventas_m2, 6=personal, 7=ticket
    ws = wb["Cuadro 6."]
    rows_out = []
    for row in ws.iter_rows(min_row=7, values_only=True):
        if not isinstance(row[0], datetime):
            continue
        rows_out.append({
            "fecha": row[0].strftime("%Y-%m"),
            "bocas": int(row[2]) if row[2] is not None else None,
            "ventas_m2": to_float(row[5]),
            "personal": int(row[6]) if row[6] is not None else None,
            "ticket_promedio": to_float(row[7]),
        })
    return rows_out


def guardar_csv_generico(rows, filepath, cols):
    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    print(f"✅ CSV guardado: {filepath} ({len(rows)} filas)")


def calcular_variacion_yoy(series, campo):
    """Calcula var_yoy para un campo dado comparando con 12 meses atrás."""
    idx = {r["fecha"]: r for r in series}
    for r in series:
        anio, mes = r["fecha"].split("-")
        fecha_ant = f"{int(anio)-1}-{mes}"
        curr = r.get(campo)
        prev = idx.get(fecha_ant, {}).get(campo)
        if curr is not None and prev is not None and prev != 0:
            r[f"{campo}_yoy"] = round((curr - prev) / prev * 100, 2)
        else:
            r[f"{campo}_yoy"] = None
    return series


def actualizar_consumo_json(cuadro1, cuadro4, cuadro6):
    ultimo_c1 = next((r for r in reversed(cuadro1) if r["indice_orig"] is not None), None)
    ultimo_c4 = next((r for r in reversed(cuadro4) if r["total_mill"] is not None), None)
    ultimo_c6 = next((r for r in reversed(cuadro6) if r["ticket_promedio"] is not None), None)

    consumo_data = {}
    if OUT_JSON.exists():
        with open(OUT_JSON, "r", encoding="utf-8") as f:
            consumo_data = json.load(f)

    fecha = ultimo_c1["fecha"] if ultimo_c1 else "—"
    consumo_data["updatedAt"] = date.today().isoformat()
    consumo_data["fuente"] = "INDEC — Encuesta de Supermercados"
    consumo_data["periodo"] = fecha

    # === Summary cards ===
    summary = consumo_data.get("summary", [])
    updates = {
        "ventas_minoristas": {
            "value": ultimo_c1["var_yoy"] if ultimo_c1 else None,
            "unit": "%",
            "foot": f"Var. interanual constantes — {fecha}",
        },
        "var_acumulada": {
            "value": ultimo_c1["var_acum"] if ultimo_c1 else None,
            "unit": "%",
            "foot": f"Acumulada {fecha[:4]} — constantes",
        },
        "var_intermensual": {
            "value": ultimo_c1["var_mom_desest"] if ultimo_c1 else None,
            "unit": "%",
            "foot": f"Var. intermensual desest. — {fecha}",
        },
        "ticket_promedio": {
            "value": round(ultimo_c6["ticket_promedio"]) if ultimo_c6 and ultimo_c6["ticket_promedio"] else None,
            "unit": "",
            "foot": f"Ticket promedio pesos corrientes — {fecha}",
        },
    }

    existing_ids = {item["id"] for item in summary}
    for sid, vals in updates.items():
        if sid in existing_ids:
            for item in summary:
                if item["id"] == sid:
                    item.update(vals)
        else:
            summary.append({"id": sid, "label": sid.replace("_", " ").title(), **vals})

    consumo_data["summary"] = summary

    # === Desglose medios de pago y canal ===
    if ultimo_c4:
        consumo_data["mediosDePago"] = {
            "fecha": ultimo_c4["fecha"],
            "efectivo_mill": ultimo_c4["efectivo_mill"],
            "debito_mill": ultimo_c4["debito_mill"],
            "credito_mill": ultimo_c4["credito_mill"],
            "otros_mill": ultimo_c4["otros_mill"],
        }
        consumo_data["canalDeVenta"] = {
            "fecha": ultimo_c4["fecha"],
            "salon_ventas_mill": ultimo_c4["salon_ventas_mill"],
            "online_mill": ultimo_c4["online_mill"],
        }

    if ultimo_c6:
        consumo_data["bocas"] = {
            "fecha": ultimo_c6["fecha"],
            "bocas": ultimo_c6["bocas"],
            "ventas_m2": ultimo_c6["ventas_m2"],
            "personal": ultimo_c6["personal"],
            "ticket_promedio": ultimo_c6["ticket_promedio"],
        }

    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(consumo_data, f, ensure_ascii=False, indent=2)

    print(f"✅ consumo.json actualizado — período: {fecha}")


def main():
    content = descargar()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    cuadro1 = procesar_cuadro1(wb)
    cuadro4 = procesar_cuadro4(wb)
    cuadro6 = procesar_cuadro6(wb)

    cuadro4 = calcular_variacion_yoy(cuadro4, "total_mill")
    cuadro4 = calcular_variacion_yoy(cuadro4, "salon_ventas_mill")
    cuadro4 = calcular_variacion_yoy(cuadro4, "online_mill")
    cuadro4 = calcular_variacion_yoy(cuadro4, "efectivo_mill")
    cuadro4 = calcular_variacion_yoy(cuadro4, "debito_mill")
    cuadro4 = calcular_variacion_yoy(cuadro4, "credito_mill")
    cuadro4 = calcular_variacion_yoy(cuadro4, "otros_mill")
    cuadro6 = calcular_variacion_yoy(cuadro6, "ticket_promedio")
    cuadro6 = calcular_variacion_yoy(cuadro6, "personal")

    guardar_csv_generico(
        cuadro1, OUT_CSV,
        ["fecha", "indice_orig", "var_yoy", "var_acum", "indice_desest", "var_mom_desest", "indice_tend", "var_mom_tend"]
    )
    guardar_csv_generico(
        cuadro4, OUT_CSV_CANALES,
        ["fecha", "total_mill", "salon_ventas_mill", "online_mill", "efectivo_mill", "debito_mill", "credito_mill", "otros_mill",
         "total_mill_yoy", "salon_ventas_mill_yoy", "online_mill_yoy", "efectivo_mill_yoy", "debito_mill_yoy", "credito_mill_yoy", "otros_mill_yoy"]
    )
    guardar_csv_generico(
        cuadro6, OUT_CSV_BOCAS,
        ["fecha", "bocas", "ventas_m2", "personal", "ticket_promedio", "ticket_promedio_yoy", "personal_yoy"]
    )

    actualizar_consumo_json(cuadro1, cuadro4, cuadro6)
    print("✅ Supermercados actualizado.")


if __name__ == "__main__":
    main()
